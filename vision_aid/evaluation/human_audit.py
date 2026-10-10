"""Preserve Human findings and select them for one page or every listed page."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from urllib.parse import urlsplit

import openpyxl

from vision_aid.evaluation.schemas import HumanAudit, HumanFinding

SHEET = "Defect report"
HEADERS = (
    "Sr. #",
    "element name",
    "Browser Combination",
    "Page name",
    "Issue Title",
    "Steps to Reproduce",
    "actual result",
    "expected result",
    "Recommendation for Fix",
    "WCAG Sc",
    "Type of Change",
    "comment",
)
ELIGIBILITY_VALUES = (
    "llm_eligible",
    "programmatic",
    "unavailable_evidence",
    "ambiguous",
)


def _source_findings(
    path: Path, *, homepage_only: bool
) -> tuple[str, tuple[HumanFinding, ...]]:
    """Read original Human finding evidence with stable source identities."""
    checksum = hashlib.sha256(path.read_bytes()).hexdigest()
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[SHEET]
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(str(value or "") for value in next(rows))
        if headers[: len(HEADERS)] != HEADERS:
            raise ValueError("Human audit has unexpected Defect report columns")
        human_findings = []
        for number, values in enumerate(rows, start=2):
            raw = {
                header: "" if value is None else str(value)
                for header, value in zip(headers, values, strict=False)
                if header
            }
            scope = raw["Page name"].strip().casefold()
            if homepage_only and scope not in {"home", "global"}:
                continue
            if not raw["Issue Title"].strip():
                raise ValueError(f"Defect report row {number} has no Issue Title")
            identity = f"{checksum}:{SHEET}:{number}".encode()
            human_findings.append(
                HumanFinding(
                    human_finding_id="human-" + hashlib.sha256(identity).hexdigest()[:16],
                    source_sheet=SHEET,
                    source_row=number,
                    page_scope=scope if scope in {"home", "global"} else "page",
                    source_element=raw["element name"],
                    problem=raw["Issue Title"],
                    raw_evidence=raw,
                    wcag_evidence=tuple(
                        dict.fromkeys(re.findall(r"\b\d+\.\d+\.\d+\b", raw["WCAG Sc"]))
                    ),
                )
            )
    finally:
        book.close()
    return checksum, tuple(human_findings)


def import_homepage_human_findings(path: Path, homepage_url: str) -> HumanAudit:
    """Preserve the one-page command's Home and Global source selection."""
    checksum, human_findings = _source_findings(path, homepage_only=True)
    if not human_findings:
        raise ValueError("Human audit has no Home or Global defects")
    return HumanAudit(path.name, checksum, homepage_url, human_findings)


def import_all_human_findings(path: Path) -> dict[str, HumanAudit]:
    """Assign local findings by reproduction URL and repeat Globals per page."""
    urls = list_audit_urls(path)
    if len(urls) != len(set(urls)):
        raise ValueError("Human audit URL list contains repeated target URLs")
    checksum, findings = _source_findings(path, homepage_only=False)
    if not findings:
        raise ValueError("Human audit has no Human findings")
    by_page: dict[str, list[HumanFinding]] = {url: [] for url in urls}
    for finding in findings:
        if finding.page_scope == "global":
            for page_findings in by_page.values():
                page_findings.append(finding)
            continue
        steps = finding.raw_evidence["Steps to Reproduce"]
        reproduced = {
            match.rstrip(".,;:)]}")
            for match in re.findall(r'https?://[^\s<>"\']+', steps)
        }
        assigned = set()
        for url in reproduced:
            if url in by_page:
                assigned.add(url)
                continue
            parsed = urlsplit(url)
            # This is the audit's known alias, not general path case folding.
            if parsed.path == "/Contact/":
                url = parsed._replace(path="/contact/").geturl()
            if url in by_page:
                assigned.add(url)
        if len(assigned) != 1:
            raise ValueError(
                f"Defect report row {finding.source_row} must map to exactly "
                "one listed target URL in Steps to Reproduce"
            )
        by_page[assigned.pop()].append(finding)
    return {
        url: HumanAudit(path.name, checksum, url, tuple(page_findings))
        for url, page_findings in by_page.items()
    }


def list_audit_urls(path: Path) -> tuple[str, ...]:
    """Read the Human audit's explicit target URLs in source order."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["URL list"]
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(str(value or "").strip() for value in next(rows))
        if "URL" not in headers:
            raise ValueError("Human audit URL list has no URL column")
        index = headers.index("URL")
        urls = tuple(
            str(row[index]).strip()
            for row in rows
            if len(row) > index and row[index] is not None and str(row[index]).strip()
        )
    finally:
        book.close()
    if not urls:
        raise ValueError("Human audit URL list has no target URLs")
    return urls
