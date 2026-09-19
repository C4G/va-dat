"""Candidate generation and reviewed one-to-one match validation."""

from __future__ import annotations

import json
from collections import Counter
from collections.abc import Sequence
from pathlib import Path
from typing import Any, Protocol

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

from vision_aid.evaluation.schemas import (
    CanonicalFinding,
    MatchDecision,
    ReferenceDefect,
    ReviewProvenance,
)


class MatchValidationError(ValueError):
    """Match decisions are unresolved, incompatible, or conflicting."""


class MatchReviewer(Protocol):
    """Adapter contract for human or future model match review."""

    def review(
        self,
        references: Sequence[ReferenceDefect],
        findings: Sequence[CanonicalFinding],
    ) -> tuple[MatchDecision, ...]:
        """Return validated decisions for candidate finding/reference pairs."""


MATCH_COLUMNS = (
    "reference_id",
    "finding_id",
    "reference_scope",
    "reference_problem",
    "reference_location",
    "reference_wcag",
    "reference_raw_evidence",
    "finding_prompt",
    "finding_problem",
    "finding_element",
    "finding_location",
    "finding_wcag",
    "finding_raw_source",
    "decision",
    "page_compatible",
    "failure_compatible",
    "location_compatible",
    "required_subdefect",
    "rationale",
    "reviewer",
    "confidence",
    "timestamp",
    "notes",
)


def generate_candidates(
    references: Sequence[ReferenceDefect],
    findings: Sequence[CanonicalFinding],
) -> tuple[tuple[str, str], ...]:
    """Remove only hard page/element incompatibilities; never award credit."""
    candidates: list[tuple[str, str]] = []
    for reference in references:
        for finding in findings:
            page_compatible = (
                reference.page_scope == "global"
                or reference.canonical_url == finding.page_url
            )
            if not page_compatible:
                continue
            reference_location = reference.location.casefold()
            finding_evidence = f"{finding.element} {finding.location}".casefold()
            if reference_location and finding_evidence:
                # Only treat explicit, disjoint element categories as impossible.
                categories = ("image", "heading", "link", "form", "table", "media")
                ref_categories = {
                    item for item in categories if item in reference_location
                }
                finding_categories = {
                    item for item in categories if item in finding_evidence
                }
                if (
                    ref_categories
                    and finding_categories
                    and ref_categories.isdisjoint(finding_categories)
                ):
                    continue
            candidates.append((reference.reference_id, finding.finding_id))
    return tuple(candidates)


def validate_match_decisions(
    references: Sequence[ReferenceDefect],
    findings: Sequence[CanonicalFinding],
    decisions: Sequence[MatchDecision],
    *,
    require_resolved: bool = True,
) -> tuple[MatchDecision, ...]:
    """Require resolved evidence-backed decisions and one-to-one findings."""
    reference_by_id = {item.reference_id: item for item in references}
    finding_ids = {item.finding_id for item in findings}
    pair_counts = Counter((item.reference_id, item.finding_id) for item in decisions)
    repeated = [pair for pair, count in pair_counts.items() if count > 1]
    if repeated:
        raise MatchValidationError(
            f"Match review repeats candidate pair {repeated[0]!r}."
        )
    for decision in decisions:
        if decision.reference_id not in reference_by_id:
            raise MatchValidationError(
                f"Unknown reference in match review: {decision.reference_id}."
            )
        if decision.finding_id not in finding_ids:
            raise MatchValidationError(
                f"Unknown finding in match review: {decision.finding_id}."
            )
        if require_resolved and decision.state == "needs_review":
            raise MatchValidationError(
                "Unresolved needs_review match decisions cannot be scored."
            )
        if not decision.provenance.reviewer or not decision.provenance.rationale:
            raise MatchValidationError("Every match decision requires provenance.")
        if not 0 <= decision.provenance.confidence <= 1:
            raise MatchValidationError("Match confidence must be 0 through 1.")
        if decision.state != "accepted":
            continue
        if not all(
            (
                decision.page_compatible,
                decision.failure_compatible,
                decision.location_compatible,
            )
        ):
            raise MatchValidationError(
                "An accepted match requires compatible page, underlying failure, "
                "and location or element-group evidence."
            )
        reference = reference_by_id[decision.reference_id]
        if reference.required_subdefects:
            if decision.required_subdefect not in reference.required_subdefects:
                raise MatchValidationError(
                    f"Accepted match for {reference.reference_id} must identify one "
                    "of its required sub-defects."
                )
        elif decision.required_subdefect:
            raise MatchValidationError(
                f"Reference {reference.reference_id} has no required sub-defects."
            )
    accepted = [item for item in decisions if item.state == "accepted"]
    finding_counts = Counter(item.finding_id for item in accepted)
    conflicts = [item for item, count in finding_counts.items() if count > 1]
    if conflicts:
        raise MatchValidationError(
            "Accepted matches must be one-to-one; finding "
            f"{conflicts[0]} claims multiple reference defects."
        )
    plain_reference_counts = Counter(
        item.reference_id
        for item in accepted
        if not reference_by_id[item.reference_id].required_subdefects
    )
    conflicts = [item for item, count in plain_reference_counts.items() if count > 1]
    if conflicts:
        raise MatchValidationError(
            "Accepted matches must be one-to-one; reference "
            f"{conflicts[0]} is claimed by multiple findings."
        )
    subdefect_counts = Counter(
        (item.reference_id, item.required_subdefect)
        for item in accepted
        if item.required_subdefect is not None
    )
    subdefect_conflicts = [
        item for item, count in subdefect_counts.items() if count > 1
    ]
    if subdefect_conflicts:
        raise MatchValidationError(
            "Accepted matches must be one-to-one at the required sub-defect level."
        )
    return tuple(decisions)


def export_match_workbook(
    references: Sequence[ReferenceDefect],
    findings: Sequence[CanonicalFinding],
    output_path: Path,
) -> None:
    """Create an evidence-rich private review workbook for plausible pairs."""
    reference_by_id = {item.reference_id: item for item in references}
    finding_by_id = {item.finding_id: item for item in findings}
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Matches"
    sheet.append(MATCH_COLUMNS)
    for reference_id, finding_id in generate_candidates(references, findings):
        reference = reference_by_id[reference_id]
        finding = finding_by_id[finding_id]
        sheet.append(
            (
                reference_id,
                finding_id,
                reference.page_scope,
                reference.problem,
                reference.location,
                ", ".join(reference.wcag_evidence),
                json.dumps(reference.raw_evidence, ensure_ascii=False, sort_keys=True),
                finding.prompt,
                finding.problem,
                finding.element,
                finding.location,
                ", ".join(finding.wcag_evidence),
                json.dumps(finding.raw_source, ensure_ascii=False, sort_keys=True),
                "needs_review",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
            )
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in ("D", "E", "G", "I", "K", "M", "S", "W"):
        sheet.column_dimensions[column].width = 40
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = openpyxl.styles.Alignment(
                wrap_text=True,
                vertical="top",
            )
    decisions = DataValidation(
        type="list",
        formula1='"accepted,rejected,needs_review"',
    )
    booleans = DataValidation(type="list", formula1='"TRUE,FALSE"')
    sheet.add_data_validation(decisions)
    sheet.add_data_validation(booleans)
    decisions.add(f"N2:N{max(2, sheet.max_row)}")
    booleans.add(f"O2:Q{max(2, sheet.max_row)}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _boolean(value: Any, *, row: int, field: str) -> bool | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().casefold()
    if normalized in {"true", "yes", "1"}:
        return True
    if normalized in {"false", "no", "0"}:
        return False
    raise MatchValidationError(f"Match row {row} has invalid {field} value.")


class HumanMatchReviewer:
    """Load human match decisions through the shared reviewer contract."""

    def __init__(self, review_path: Path):
        self.review_path = review_path

    def review(
        self,
        references: Sequence[ReferenceDefect],
        findings: Sequence[CanonicalFinding],
    ) -> tuple[MatchDecision, ...]:
        """Load provenance-complete decisions and validate accepted conflicts."""
        try:
            workbook = openpyxl.load_workbook(
                self.review_path,
                read_only=True,
                data_only=True,
            )
            sheet = workbook["Matches"]
        except (OSError, KeyError, ValueError) as error:
            raise MatchValidationError(
                f"Could not read Matches review workbook: {error}"
            ) from error
        rows = sheet.iter_rows(values_only=True)
        header_values = next(rows, None) or ()
        headers = [_text(value) for value in header_values]
        missing = set(MATCH_COLUMNS) - set(headers)
        if missing:
            raise MatchValidationError(
                "Matches review is missing columns: " + ", ".join(sorted(missing))
            )
        decisions = []
        for row_number, values in enumerate(rows, start=2):
            record = dict(zip(headers, values, strict=False))
            reference_id = _text(record.get("reference_id"))
            finding_id = _text(record.get("finding_id"))
            if not reference_id and not finding_id:
                continue
            state = _text(record.get("decision"))
            if state not in {"accepted", "rejected", "needs_review"}:
                raise MatchValidationError(
                    f"Match row {row_number} has invalid decision {state!r}."
                )
            reviewer = _text(record.get("reviewer"))
            rationale = _text(record.get("rationale"))
            timestamp = _text(record.get("timestamp"))
            if not reviewer or not rationale or not timestamp:
                raise MatchValidationError(
                    f"Match row {row_number} requires reviewer, rationale, and timestamp."
                )
            try:
                raw_confidence = record.get("confidence")
                if raw_confidence is None:
                    raise ValueError
                confidence = float(raw_confidence)
            except (TypeError, ValueError) as error:
                raise MatchValidationError(
                    f"Match row {row_number} requires numeric confidence."
                ) from error
            decisions.append(
                MatchDecision(
                    reference_id=reference_id,
                    finding_id=finding_id,
                    state=state,  # type: ignore[arg-type]
                    provenance=ReviewProvenance(
                        reviewer=reviewer,
                        rationale=rationale,
                        confidence=confidence,
                        timestamp=timestamp,
                    ),
                    required_subdefect=(
                        _text(record.get("required_subdefect")) or None
                    ),
                    page_compatible=_boolean(
                        record.get("page_compatible"),
                        row=row_number,
                        field="page_compatible",
                    ),
                    failure_compatible=_boolean(
                        record.get("failure_compatible"),
                        row=row_number,
                        field="failure_compatible",
                    ),
                    location_compatible=_boolean(
                        record.get("location_compatible"),
                        row=row_number,
                        field="location_compatible",
                    ),
                    notes=_text(record.get("notes")),
                )
            )
        workbook.close()
        return validate_match_decisions(
            references,
            findings,
            decisions,
            require_resolved=False,
        )
