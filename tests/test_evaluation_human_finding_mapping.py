"""Exercise approved Human audit decisions through the evaluator CLI."""

import csv
import json

import pytest
from test_evaluation import (
    edit_review,
    fake_client,
    load,
    only_run,
    refused,
    report,
    run_cli,
)
from test_evaluation import workspace as workspace  # noqa: PLC0414 - expose shared pytest fixture

URL = "https://example.test/"


@pytest.fixture
def approved(workspace):
    """Author independent page-level expectations from a synthetic Human audit."""
    (workspace / "home.html").write_text(
        '<html><head></head><body><h2>Home</h2><a href="/">Click here</a></body></html>'
    )
    assert run_cli(workspace) == 0
    seed = only_run(workspace)
    human = load(seed, "artifacts/human-findings.json")
    classifications = [
        "programmatic", "llm_eligible", "llm_eligible",
        "unavailable_evidence", "ambiguous", "programmatic",
    ]
    expectations = [
        [
            {"rule": "PAGE_TITLE_001", "evidence": {
                "raw_source": {"location": {"tag": "head"}}}},
            {"rule": "LANG_001", "evidence": {
                "problem": "<html> element missing lang attribute."}},
        ],
        [{"rule": "LAND_001", "evidence": {
            "problem": "Page does not contain a <main> landmark."}}],
        [], [], [],
        [{"rule": "HEAD_003", "evidence": {
            "problem": "The page does not contain an <h1> element."}}],
    ]
    mapping = {
        "schema_version": 1,
        "human_audit_sha256": human["human_audit_sha256"],
        "approved_by": "Independent synthetic audit review",
        "decisions": [
            {
                "human_finding_id": finding["human_finding_id"],
                "classification": classification,
                "classification_reason": f"Approved responsibility for row {index + 2}",
                "programmatic_by_page": {URL: {
                    "expectations": matches,
                    "reason": "Complete page defect" if matches else "No existing rule catches this defect",
                }},
            }
            for index, (finding, classification, matches) in enumerate(zip(
                human["human_findings"], classifications, expectations, strict=True
            ))
        ],
    }
    path = workspace / "human-finding-mapping.json"
    path.write_text(json.dumps(mapping))
    return path, seed


def review_rows(run):
    """Read the operator-facing review rows in original audit order."""
    with (run / "review.csv").open(newline="") as stream:
        return sorted(csv.DictReader(stream), key=lambda row: int(row["source_row"]))


def human_finding_mapping_run(workspace, path, *extra):
    """Prepare a mapping-backed run through the normal command."""
    root = workspace / ".model-evaluation/runs"
    before = set(root.iterdir())
    assert run_cli(workspace, "--human-finding-mapping", str(path), *extra) == 0
    (run,) = set(root.iterdir()) - before
    return run


def test_preparation_populates_approved_decisions_and_preserves_human_evidence(
    workspace, approved, monkeypatch,
):
    """Preview resolves joint and cross-catches and keeps source evidence intact."""
    path, seed = approved
    monkeypatch.setattr("anthropic.Anthropic", lambda **_: pytest.fail("paid request"))
    run = human_finding_mapping_run(workspace, path)
    original = review_rows(seed)
    rows = review_rows(run)
    assert [row["classification"] for row in rows] == [
        "programmatic", "llm_eligible", "llm_eligible",
        "unavailable_evidence", "ambiguous", "programmatic",
    ]
    findings = {
        item["finding_id"]: item
        for item in load(run, "normalized-programmatic-findings.json")
    }
    title_id, lang_id = rows[0]["programmatic_finding_id"].split(";")
    assert findings[title_id]["problem_family"] == "PAGE_TITLE_001"
    assert findings[title_id]["raw_source"]["location"]["tag"] == "head"
    assert findings[lang_id]["problem_family"] == "LANG_001"
    assert rows[1]["programmatic_finding_id"]
    assert all(not row["programmatic_finding_id"] for row in rows[2:5])
    assert rows[0]["programmatic_match_reason"] == "Complete page defect"
    for previous, current in zip(original, rows, strict=True):
        for field in ("human_finding_id", "source_sheet", "source_row", "page_scope",
                      "source_element", "actual result", "Issue Title"):
            assert current[field] == previous[field]
    assert load(run, "artifacts/human-findings.json") == load(seed, "artifacts/human-findings.json")
    assert load(run, "artifacts/evaluation-baseline.json") == json.loads(path.read_text())


def test_models_share_fixed_decisions_but_keep_independent_llm_review(
    workspace, approved, monkeypatch,
):
    """Each model resolves its own IDs; only independent LLM decisions affect its score."""
    path, _ = approved
    fake_client(monkeypatch,
        {"success": True, "response": '{"issues": ["Missing h1"]}', "usage": {}},
        {"success": True, "response": '[{"text": "Click here", "is_clear": false, "issues": ["Vague link"]}]', "usage": {}},
    )
    first = human_finding_mapping_run(workspace, path, "--live", "--approve-live")
    fake_client(monkeypatch)
    second = human_finding_mapping_run(workspace, path, "--model", "claude-sonnet-5-5", "--live", "--approve-live")
    llm = load(first, "normalized-llm-findings.json")
    assert len(llm) == 2
    edit_review(first,
        {"llm_finding_id": llm[0]["finding_id"], "match_reason": "Same defect despite remediation wording"},
        {},
        {"llm_finding_id": llm[1]["finding_id"], "match_reason": "Full semantic match", "review_notes": "Different example remains valid"},
    )
    fixed = ["classification", "classification_reason", "programmatic_match_reason"]
    first_rows, second_rows = review_rows(first), review_rows(second)
    by_run = [
        {item["finding_id"]: item["problem_family"] for item in load(run, "normalized-programmatic-findings.json")}
        for run in (first, second)
    ]
    for a, b in zip(first_rows, second_rows, strict=True):
        assert [a[key] for key in fixed] == [b[key] for key in fixed]
        assert len(a["programmatic_finding_id"].split(";")) == len(b["programmatic_finding_id"].split(";"))
        assert {by_run[0][identity] for identity in a["programmatic_finding_id"].split(";") if identity} == {
            by_run[1][identity] for identity in b["programmatic_finding_id"].split(";") if identity
        }
        if a["programmatic_finding_id"]:
            assert a["programmatic_finding_id"] != b["programmatic_finding_id"]
    # Later deliberate shared edits must not change either saved interpretation.
    mapping = json.loads(path.read_text())
    mapping["decisions"][0]["classification_reason"] = "A later manual update"
    path.write_text(json.dumps(mapping))
    first_report, second_report = report(first), report(second)
    assert "Independent synthetic audit review" in first_report
    assert "Complete page defect" in first_report
    assert "Same defect despite remediation wording" in first_report
    assert "Different example remains valid" in first_report
    assert "Approved responsibility for row 2" in first_report
    assert "A later manual update" not in first_report
    assert "LLM: 1 caught / 2 eligible" in first_report and "LLM: 0 caught / 2 eligible" in second_report
    assert "Overall: 4 caught / 6 eligible" in first_report and "Overall: 3 caught / 6 eligible" in second_report
    assert "Programmatic: 2 caught / 2 eligible" in first_report and "Programmatic: 2 caught / 2 eligible" in second_report


@pytest.mark.parametrize("column,value", [
    ("classification", "llm_eligible"),
    ("classification_reason", "A run-specific reinterpretation"),
    ("programmatic_finding_id", ""),
    ("programmatic_match_reason", "A different justification"),
    ("actual result", "Changed original Human evidence"),
])
def test_reporting_rejects_edited_approved_decisions_and_removes_stale_report(
    workspace, approved, monkeypatch, capsys, column, value,
):
    """Editable review cells cannot rewrite fixed decisions or original evidence."""
    path, _ = approved
    fake_client(monkeypatch)
    run = human_finding_mapping_run(workspace, path, "--live", "--approve-live")
    report(run)
    edit_review(run, {column: value})
    message = refused(run, capsys)
    assert "Human finding" in message
    assert "changed" in message


@pytest.mark.parametrize("change,error", [
    ("audit", "Human audit identity differs"),
    ("decision", "missing approved decision"),
    ("reason", "needs approved classification_reason"),
    ("page", "missing explicit Programmatic page decision"),
    ("missing", "resolved to 0 findings"),
    ("reuse", "finding ID reuse"),
    ("programmatic_absence", "requires expected coverage"),
    ("not_testable", "not testable"),
])
def test_preparation_rejects_invalid_or_unresolved_approved_expectations(
    workspace, approved, capsys, change, error,
):
    """Missing decisions and evidence fail with useful finding/page context."""
    path, _ = approved
    mapping = json.loads(path.read_text())
    if change == "audit":
        mapping["human_audit_sha256"] = "unrelated-audit"
    elif change == "decision":
        mapping["decisions"].pop(0)
    elif change == "reason":
        mapping["decisions"][0]["classification_reason"] = ""
    elif change == "page":
        mapping["decisions"][0]["programmatic_by_page"]["https://elsewhere.test/"] = mapping["decisions"][0]["programmatic_by_page"].pop(URL)
    elif change == "missing":
        mapping["decisions"][0]["programmatic_by_page"][URL]["expectations"][0]["evidence"] = {"element": "unreported-element"}
    elif change == "reuse":
        mapping["decisions"][1]["programmatic_by_page"][URL]["expectations"] = mapping["decisions"][0]["programmatic_by_page"][URL]["expectations"]
    elif change == "programmatic_absence":
        mapping["decisions"][0]["programmatic_by_page"][URL]["expectations"] = []
    elif change == "not_testable":
        mapping["decisions"][0]["classification"] = "unavailable_evidence"
    path.write_text(json.dumps(mapping))
    with pytest.raises(SystemExit) as caught:
        run_cli(workspace, "--human-finding-mapping", str(path))
    assert caught.value.code == 2
    assert error in capsys.readouterr().err


@pytest.mark.parametrize("change,count", [("missing", 0), ("duplicate", 2)])
def test_reporting_requires_unique_actual_checker_evidence(
    workspace, approved, monkeypatch, capsys, change, count,
):
    """A saved expected catch cannot manufacture credit after evidence is lost."""
    path, _ = approved
    fake_client(monkeypatch)
    run = human_finding_mapping_run(workspace, path, "--live", "--approve-live")
    findings = load(run, "normalized-programmatic-findings.json")
    if change == "missing":
        findings = [item for item in findings if item["problem_family"] != "PAGE_TITLE_001"]
    else:
        duplicate = next(item.copy() for item in findings if item["problem_family"] == "PAGE_TITLE_001")
        duplicate["finding_id"] = "programmatic-second-valid-id"
        findings.append(duplicate)
    (run / "normalized-programmatic-findings.json").write_text(json.dumps(findings))
    message = refused(run, capsys)
    assert "PAGE_TITLE_001" in message
    assert "Human finding human-" in message
    assert f"resolved to {count} findings" in message


def test_reporting_rejects_llm_credit_for_not_testable_findings(
    workspace, approved, monkeypatch, capsys,
):
    """LLM edits retain existing not-testable restrictions even with approved decisions."""
    path, _ = approved
    fake_client(monkeypatch,
        {"success": True, "response": '{"issues": ["A concrete defect"]}', "usage": {}},
    )
    run = human_finding_mapping_run(workspace, path, "--live", "--approve-live")
    identity = load(run, "normalized-llm-findings.json")[0]["finding_id"]
    edit_review(run, {}, {}, {}, {"llm_finding_id": identity, "match_reason": "Claimed catch"})
    assert "not testable" in refused(run, capsys)


def test_reporting_rejects_modified_saved_approved_source(
    workspace, approved, monkeypatch, capsys,
):
    """Historical interpretation is pinned to the copy saved with the run."""
    path, _ = approved
    fake_client(monkeypatch)
    run = human_finding_mapping_run(workspace, path, "--live", "--approve-live")
    copied = run / "artifacts/evaluation-baseline.json"
    mapping = json.loads(copied.read_text())
    mapping["decisions"][0]["classification_reason"] = "Rewritten history"
    copied.write_text(json.dumps(mapping))
    assert "checksum changed" in refused(run, capsys)


def test_bundled_human_finding_mapping_is_automatic_and_cannot_be_removed_to_unlock_reporting(
    workspace, approved, monkeypatch, capsys,
):
    """Known audits automatically use approved decisions and require saved approval."""
    path, _ = approved
    mapping = json.loads(path.read_text())
    bundled = workspace / "bundled-mappings"
    bundled.mkdir()
    (bundled / f"{mapping['human_audit_sha256']}.json").write_text(path.read_text())
    monkeypatch.setattr("vision_aid.evaluation.human_finding_mapping.HUMAN_FINDING_MAPPING_DIRECTORY", bundled)
    fake_client(monkeypatch)
    before = set((workspace / ".model-evaluation/runs").iterdir())
    assert run_cli(workspace, "--live", "--approve-live") == 0
    (run,) = set((workspace / ".model-evaluation/runs").iterdir()) - before
    assert review_rows(run)[0]["classification"] == "programmatic"
    report(run)
    (run / "artifacts/evaluation-baseline.json").unlink()
    manifest = load(run, "run.json")
    manifest.pop("evaluation_baseline_sha256")
    manifest.pop("evaluation_baseline_approved_by")
    (run / "run.json").write_text(json.dumps(manifest))
    assert "prepare a new run with --human-finding-mapping PATH" in refused(run, capsys)


@pytest.mark.parametrize("change,error", [
    ("unknown", "Unknown LLM finding ID"),
    ("source", "Unknown LLM finding ID"),
    ("reuse", "Conflicting LLM finding ID"),
])
def test_human_finding_mapping_does_not_relax_llm_identity_or_reuse_validation(
    workspace, approved, monkeypatch, capsys, change, error,
):
    """Editable LLM choices must still belong to this run and one Human finding."""
    path, _ = approved
    fake_client(monkeypatch,
        {"success": True, "response": '{"issues": ["A concrete defect"]}', "usage": {}},
    )
    run = human_finding_mapping_run(workspace, path, "--live", "--approve-live")
    llm = load(run, "normalized-llm-findings.json")[0]["finding_id"]
    if change == "unknown":
        identity = "llm-another-run"
    elif change == "source":
        identity = review_rows(run)[0]["programmatic_finding_id"].split(";")[0]
    else:
        identity = llm
    edit_review(run,
        {"llm_finding_id": identity, "match_reason": "Claimed semantic match"},
        {"llm_finding_id": llm, "match_reason": "Reused match"} if change == "reuse" else {},
    )
    assert error in refused(run, capsys)


def test_single_page_uses_listed_page_local_findings_and_globals(
    workspace, capsys,
):
    """A non-home URL receives its own Human findings and approved page expectations."""
    import openpyxl

    careers = "https://example.test/careers/"
    book = openpyxl.load_workbook(workspace / "human-audit.xlsx")
    sheet = book["Defect report"]
    sheet.cell(2, 6, f"Open {URL}")
    sheet.cell(8, 6, f"Open {careers}")
    urls = book.create_sheet("URL list")
    urls.append(["URL"])
    urls.append([URL])
    urls.append([careers])
    book.save(workspace / "human-audit.xlsx")
    book.close()
    assert run_cli(workspace, "--source-url", careers) == 0
    seed = only_run(workspace)
    human = load(seed, "artifacts/human-findings.json")
    assert [finding["source_row"] for finding in human["human_findings"]] == [3, 4, 5, 6, 7, 8]
    path = workspace / "careers-human-finding-mapping.json"
    path.write_text(json.dumps({
        "schema_version": 1, "human_audit_sha256": human["human_audit_sha256"],
        "approved_by": "Non-home reviewer", "decisions": [
            {
                "human_finding_id": item["human_finding_id"],
                "classification": "llm_eligible", "classification_reason": "Model judgment required",
                "programmatic_by_page": {careers: {"expectations": [], "reason": "No complete checker catch"}},
            } for item in human["human_findings"]
        ],
    }))
    run = human_finding_mapping_run(workspace, path, "--source-url", careers)
    assert [row["source_row"] for row in review_rows(run)] == ["3", "4", "5", "6", "7", "8"]
    assert all(row["classification"] == "llm_eligible" for row in review_rows(run))
    with pytest.raises(SystemExit):
        run_cli(workspace, "--source-url", "https://example.test/unlisted/")
    assert "not listed in the Human audit" in capsys.readouterr().err
