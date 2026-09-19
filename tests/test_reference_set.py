import hashlib
from pathlib import Path

import openpyxl
import pytest

from vision_aid.evaluation.references import (
    HumanEligibilityReviewer,
    ReferenceSetError,
    export_eligibility_workbook,
    import_homepage_references,
)


def _write_workbook(path: Path) -> str:
    """Write a synthetic multi-scope workbook and return its checksum."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Defects"
    sheet.append(
        ["Page", "URL", "Problem", "Location", "WCAG", "Recommendation"]
    )
    sheet.append(
        [
            "Home",
            "https://example.test/",
            (
                "Decorative image has descriptive alt text\n"
                "(second line retained)"
            ),
            "Hero image",
            "1.1.1",
            "Use empty alternative text",
        ]
    )
    sheet.append(
        [
            "Global",
            "All pages",
            "Heading labels are unclear",
            "Primary navigation",
            "2.4.6",
            "Describe the section",
        ]
    )
    sheet.append(
        [
            "Contact",
            "https://example.test/contact",
            "Contact-only issue",
            "Form",
            "3.3.2",
            "Add instructions",
        ]
    )
    workbook.save(path)
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_import_homepage_references_retains_every_home_and_global_source_row(
    tmp_path: Path,
) -> None:
    """Home and Global imports remain lossless and source-addressed."""
    workbook_path = tmp_path / "synthetic.xlsx"
    checksum = _write_workbook(workbook_path)

    reference_set = import_homepage_references(
        workbook_path,
        expected_sha256=checksum,
        homepage_url="https://example.test/",
        version="synthetic-v1",
    )

    assert [
        (item.source_sheet, item.source_row)
        for item in reference_set.references
    ] == [
        ("Defects", 2),
        ("Defects", 3),
    ]
    assert reference_set.references[0].raw_evidence["Problem"] == (
        "Decorative image has descriptive alt text\n(second line retained)"
    )
    assert reference_set.references[1].page_scope == "global"
    assert reference_set.references[1].canonical_url == "https://example.test/"
    assert reference_set.workbook_sha256 == checksum

    with pytest.raises(ReferenceSetError, match="checksum drift"):
        import_homepage_references(
            workbook_path,
            expected_sha256="0" * 64,
            homepage_url="https://example.test/",
            version="synthetic-v1",
        )


def test_human_eligibility_review_round_trips_validated_provenance(
    tmp_path: Path,
) -> None:
    """Eligibility decisions validate and retain human provenance."""
    workbook_path = tmp_path / "synthetic.xlsx"
    checksum = _write_workbook(workbook_path)
    reference_set = import_homepage_references(
        workbook_path,
        expected_sha256=checksum,
        homepage_url="https://example.test/",
        version="synthetic-v1",
    )
    review_path = tmp_path / "eligibility.xlsx"
    export_eligibility_workbook(reference_set, review_path)

    review = openpyxl.load_workbook(review_path)
    sheet = review["Eligibility"]
    headings = {cell.value: cell.column for cell in sheet[1]}
    decisions = [
        ("llm_eligible", "accepted", "Human judgment is required."),
        ("programmatic", "accepted", "A deterministic rule owns this defect."),
    ]
    for row, (classification, state, rationale) in enumerate(
        decisions, start=2
    ):
        sheet.cell(row, headings["classification"], classification)
        sheet.cell(row, headings["decision"], state)
        sheet.cell(row, headings["rationale"], rationale)
        sheet.cell(row, headings["reviewer"], "reviewer@example.test")
        sheet.cell(row, headings["confidence"], 0.9)
        sheet.cell(row, headings["timestamp"], "2026-09-19T12:00:00Z")
    review.save(review_path)

    reviewed = HumanEligibilityReviewer(review_path).review(reference_set)

    assert [item.eligibility for item in reviewed.references] == [
        "llm_eligible",
        "programmatic",
    ]
    assert reviewed.references[0].eligibility_review is not None
    assert reviewed.references[0].eligibility_review.reviewer == (
        "reviewer@example.test"
    )

    sheet.cell(2, headings["classification"], "secret_fifth_state")
    review.save(review_path)
    with pytest.raises(ReferenceSetError, match="classification"):
        HumanEligibilityReviewer(review_path).review(reference_set)
